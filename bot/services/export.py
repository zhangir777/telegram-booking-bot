import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

COLUMNS = ["ID", "Date", "Time", "Service", "Master", "Price, KZT", "Client", "Phone", "Status"]


def booking_rows(bookings) -> list[list]:
    return [
        [
            b.id,
            b.start.strftime("%Y-%m-%d"),
            b.start.strftime("%H:%M"),
            b.service.name,
            b.master.name,
            b.service.price,
            b.user.name,
            b.user.phone or "",
            b.status,
        ]
        for b in bookings
    ]


def to_xlsx(bookings) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bookings"
    ws.append(COLUMNS)
    for row in booking_rows(bookings):
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDEFGHI", (6, 12, 7, 24, 16, 11, 22, 16, 10), strict=True):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(bookings) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(COLUMNS)
    writer.writerows(booking_rows(bookings))
    return buf.getvalue().encode("utf-8-sig")
